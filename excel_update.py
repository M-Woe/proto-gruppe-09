"""
Schreibe die Querschnittsdaten als echte Zahlen in die Excel-Tabelle.
"""
import openpyxl, json
from openpyxl.utils import get_column_letter

xlsx_path = "data/Aufgabe 1 Excel-Liste.xlsx"
out_path = "data/Aufgabe 1 Excel-Liste aktualisiert.xlsx"

with open("data/staebe_mit_querschnitt.json", "r") as f:
    data = json.load(f)

wb = openpyxl.load_workbook(xlsx_path)
if "staebe_mit_querschnitt" in wb.sheetnames:
    del wb["staebe_mit_querschnitt"]
ws = wb.create_sheet("staebe_mit_querschnitt")

headers = list(data[0].keys())
numeric_cols = {"laenge_m","dicke_m","breite1_m","breite2_m",
                "A1_m2","A2_m2","A_avg_m2","volumen_m3"}

for c, h in enumerate(headers, 1):
    ws.cell(1, c, h)

for ri, row in enumerate(data, 2):
    for ci, h in enumerate(headers, 1):
        if h in numeric_cols:
            ws.cell(ri, ci, float(row[h]))
        else:
            ws.cell(ri, ci, str(row[h]))

for ci, h in enumerate(headers, 1):
    cl = get_column_letter(ci)
    width = max(10, len(h)+2, 15 if h in numeric_cols else 12)
    ws.column_dimensions[cl].width = width

# Zahlenformat
for ci, h in enumerate(headers, 1):
    if h in numeric_cols:
        for ri in range(2, len(data)+2):
            ws.cell(ri, ci).number_format = '0.00000000'

wb.save(out_path)
print(f"Excel mit Zahlen gespeichert: {out_path}")
print(f"  Sheet 'staebe_mit_querschnitt' ({len(data)} Zeilen)")

# Test: eine Zeile zuruecklesen
wb2 = openpyxl.load_workbook(out_path)
ws2 = wb2["staebe_mit_querschnitt"]
print(f"\nTest Flat Bar 225/20 (Zeile 2):")
for c in range(1, len(headers)+1):
    h = ws2.cell(1, c).value
    v = ws2.cell(2, c).value
    print(f"  {h}: {v}  (Typ: {type(v).__name__})")