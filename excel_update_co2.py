"""Fuege CO2-Sheet in Excel ein."""
import openpyxl, json
from openpyxl.utils import get_column_letter

xlsx_path = "data/Aufgabe 1 Excel-Liste.xlsx"
with open("data/staebe_mit_co2.json") as f:
    data = json.load(f)

wb = openpyxl.load_workbook(xlsx_path)
if "staebe_mit_co2" in wb.sheetnames:
    del wb["staebe_mit_co2"]

ws = wb.create_sheet("staebe_mit_co2")
headers = list(data[0].keys())
num = {"volumen_m3","masse_kg","laenge_m","co2_A1A3_kg","co2_C_sum_kg","co2_D_kg","bewehrung_kg","bewehrung_co2_A1A3_kg"}

for c, h in enumerate(headers, 1):
    ws.cell(1, c, h)
for ri, row in enumerate(data, 2):
    for ci, h in enumerate(headers, 1):
        v = row.get(h, "")
        if h in num and v != "":
            try: ws.cell(ri, ci, float(v)); continue
            except: pass
        ws.cell(ri, ci, str(v))

for ci, h in enumerate(headers, 1):
    ws.column_dimensions[get_column_letter(ci)].width = max(10, len(h)+2)
for ci, h in enumerate(headers, 1):
    if h in num:
        for ri in range(2, len(data)+2):
            ws.cell(ri, ci).number_format = "0.000000"

wb.save(xlsx_path)
print(f"Excel aktualisiert: {xlsx_path}")
print(f"  Sheet 'staebe_mit_co2' ({len(data)} Zeilen)")