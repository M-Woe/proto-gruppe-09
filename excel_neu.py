"""
Excel-Tabelle NEU erstellt. Alle Sheets mit korrekten float-Werten.
"""
import openpyxl, json, csv
from openpyxl.utils import get_column_letter
from collections import defaultdict

with open("data/staebe_mit_querschnitt.json") as f:
    qs_data = json.load(f)
with open("data/staebe_linecurves.csv", encoding="utf-8-sig") as f:
    lc_rows = list(csv.DictReader(f))
with open("data/staebe_breps.csv", encoding="utf-8-sig") as f:
    br_rows = list(csv.DictReader(f))

wb = openpyxl.Workbook()

def write_sheet(ws, headers, rows, numeric=None):
    if numeric is None:
        numeric = set()
    for c, h in enumerate(headers, 1):
        ws.cell(1, c, h)
    for ri, row in enumerate(rows, 2):
        for ci, h in enumerate(headers, 1):
            v = row.get(h, "")
            if h in numeric and v != "":
                try:
                    ws.cell(ri, ci, float(v))
                    continue
                except:
                    pass
            ws.cell(ri, ci, str(v))
    for ci, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(ci)].width = max(10, len(h) + 2)

# === Sheet 1: staebe_linecurves ===
ws1 = wb.active
ws1.title = "staebe_linecurves"
lc_num = {"from_x","from_y","from_z","to_x","to_y","to_z","length_m","obj_index"}
lc_hdrs = list(lc_rows[0].keys())
write_sheet(ws1, lc_hdrs, lc_rows, lc_num)
for ri in range(2, len(lc_rows) + 2):
    for ci in range(1, len(lc_hdrs) + 1):
        if lc_hdrs[ci - 1] in lc_num:
            ws1.cell(ri, ci).number_format = "0.0000"

# === Sheet 2: staebe_breps (korrigierte Querschnittswerte) ===
ws2 = wb.create_sheet("staebe_breps")
qs_lookup = {r["br_uuid"]: r for r in qs_data}
br_hdrs = ["uuid","layer","name","dicke_m","breite1_m","breite2_m",
           "A1_m2","A2_m2","A_avg_m2","is_solid","obj_index"]
br_num = {"dicke_m","breite1_m","breite2_m","A1_m2","A2_m2","A_avg_m2","obj_index"}
br_out = []
for row in br_rows:
    uid = row.get("uuid","")
    qr = qs_lookup.get(uid)
    br_out.append({
        "uuid": uid, "layer": row.get("layer",""), "name": row.get("name",""),
        "dicke_m": float(qr["dicke_m"]) if qr else 0,
        "breite1_m": float(qr["breite1_m"]) if qr else 0,
        "breite2_m": float(qr["breite2_m"]) if qr else 0,
        "A1_m2": float(qr["A1_m2"]) if qr else 0,
        "A2_m2": float(qr["A2_m2"]) if qr else 0,
        "A_avg_m2": float(qr["A_avg_m2"]) if qr else 0,
        "is_solid": "True",
        "obj_index": int(row.get("obj_index",0)),
    })
write_sheet(ws2, br_hdrs, br_out, br_num)
for ci, h in enumerate(br_hdrs, 1):
    if h in br_num:
        for ri in range(2, len(br_out) + 2):
            ws2.cell(ri, ci).number_format = "0.00000000"

# === Sheet 3: staebe_layer_tabelle ===
ws3 = wb.create_sheet("staebe_layer_tabelle")
tab_hdrs = ["Layer","Anzahl","Profil","Kategorie"]
tab_rows = [
    {"Layer":"Flat Bar 33.3/20","Anzahl":116,"Profil":"Flachstab 33,3x20mm","Kategorie":"Flachstab"},
    {"Layer":"Flat Bar 275/12","Anzahl":70,"Profil":"Flachstab 275x12mm","Kategorie":"Flachstab"},
    {"Layer":"Flat Bar 275/20","Anzahl":41,"Profil":"Flachstab 275x20mm","Kategorie":"Flachstab"},
    {"Layer":"Flat Bar 162.5/12","Anzahl":35,"Profil":"Flachstab 162,5x12mm","Kategorie":"Flachstab"},
    {"Layer":"Flat Bar 175/12","Anzahl":35,"Profil":"Flachstab 175x12mm","Kategorie":"Flachstab"},
    {"Layer":"Flat Bar 200/12","Anzahl":35,"Profil":"Flachstab 200x12mm","Kategorie":"Flachstab"},
    {"Layer":"Flat Bar 225/12","Anzahl":35,"Profil":"Flachstab 225x12mm","Kategorie":"Flachstab"},
    {"Layer":"Flat Bar 237.5/12","Anzahl":35,"Profil":"Flachstab 237,5x12mm","Kategorie":"Flachstab"},
    {"Layer":"Flat Bar 162.5/20","Anzahl":18,"Profil":"Flachstab 162,5x20mm","Kategorie":"Flachstab"},
    {"Layer":"Flat Bar 200/20","Anzahl":18,"Profil":"Flachstab 200x20mm","Kategorie":"Flachstab"},
    {"Layer":"Flat Bar 237.5/20","Anzahl":18,"Profil":"Flachstab 237,5x20mm","Kategorie":"Flachstab"},
    {"Layer":"Flat Bar 175/20","Anzahl":23,"Profil":"Flachstab 175x20mm","Kategorie":"Flachstab"},
    {"Layer":"Flat Bar 225/20","Anzahl":23,"Profil":"Flachstab 225x20mm","Kategorie":"Flachstab"},
    {"Layer":"RO 139.7x8.0 | EN 10210-2:2006","Anzahl":12,"Profil":"Rundrohr 139,7x8mm","Kategorie":"Rundrohr"},
    {"Layer":"Rectangle 500/500","Anzahl":47,"Profil":"Kastenprofil 500x500mm","Kategorie":"Kastenprofil"},
]
write_sheet(ws3, tab_hdrs, tab_rows, {"Anzahl"})

# === Sheet 4: staebe_mit_querschnitt ===
ws4 = wb.create_sheet("staebe_mit_querschnitt")
qs_hdrs = list(qs_data[0].keys())
qs_num = {"laenge_m","dicke_m","breite1_m","breite2_m","A1_m2","A2_m2","A_avg_m2","volumen_m3"}
write_sheet(ws4, qs_hdrs, qs_data, qs_num)
for ci, h in enumerate(qs_hdrs, 1):
    if h in qs_num:
        for ri in range(2, len(qs_data) + 2):
            ws4.cell(ri, ci).number_format = "0.00000000"

# === Sheet 5: Zusammenfassung ===
ws5 = wb.create_sheet("Zusammenfassung")
by_ly = defaultdict(lambda: {"n":0,"V":0,"L":0,"Ac":0})
for r in qs_data:
    b = by_ly[r["layer"]]
    b["n"] += 1
    b["V"] += float(r["volumen_m3"])
    b["L"] += float(r["laenge_m"])
    b["Ac"] += float(r["A_avg_m2"])
s5_hdrs = ["Layer","Anzahl","Laenge_m","Volumen_m3","A_avg_m2"]
for c, h in enumerate(s5_hdrs, 1):
    ws5.cell(1, c, h)
ri = 2
for ln in sorted(by_ly):
    b = by_ly[ln]
    ws5.cell(ri, 1, ln)
    ws5.cell(ri, 2, b["n"])
    ws5.cell(ri, 3, round(b["L"], 3))
    ws5.cell(ri, 4, round(b["V"], 6))
    ws5.cell(ri, 5, round(b["Ac"] / b["n"], 8))
    ri += 1
ws5.cell(ri, 1, "SUMME")
ws5.cell(ri, 2, sum(b["n"] for b in by_ly.values()))
ws5.cell(ri, 3, round(sum(b["L"] for b in by_ly.values()), 3))
ws5.cell(ri, 4, round(sum(b["V"] for b in by_ly.values()), 6))
for ci in range(1, 6):
    ws5.column_dimensions[get_column_letter(ci)].width = 18

# === Speichern ===
wb.save("data/Aufgabe 1 Excel-Liste.xlsx")
print(f"Sheets: {wb.sheetnames}")

# Test: eine Zeile zuruecklesen
wb2 = openpyxl.load_workbook("data/Aufgabe 1 Excel-Liste.xlsx")
for sn in wb2.sheetnames:
    ws = wb2[sn]
    print(f"  {sn}: {ws.max_row} Zeilen x {ws.max_column} Spalten")
    # Printe Header + erste Datenzeile
    if ws.max_row >= 2:
        hdrs = [ws.cell(1,c).value for c in range(1, ws.max_column+1)]
        first = [ws.cell(2,c).value for c in range(1, ws.max_column+1)]
        print(f"    Header: {hdrs}")
        print(f"    Erste:  {first}")