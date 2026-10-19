"""Web-App: 3D-Ansicht des Rhino-Modells + Staebe-Tabelle (Auswahl im Modell markiert die Zeile).

Start:  python -m uvicorn app.main:app --reload      ->  http://127.0.0.1:8000
"""
from __future__ import annotations

import io
from pathlib import Path

from fastapi import FastAPI
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles

from .model import load_model

STATIC = Path(__file__).parent / "static"

app = FastAPI(title="Stockerau Staebe-Viewer")
app.mount("/static", StaticFiles(directory=STATIC), name="static")


@app.get("/")
def index() -> FileResponse:
    return FileResponse(STATIC / "index.html")


@app.get("/api/members")
def members() -> dict:
    """Tabellendaten ohne Geometrie (leicht)."""
    data = load_model()
    light = [{k: v for k, v in m.items() if k not in ("mesh", "line")} for m in data["members"]]
    return {
        "source": data["source"],
        "member_count": data["member_count"],
        "sections": data["sections"],
        "materials": data["materials"],
        "members": light,
    }


@app.get("/api/geometry")
def geometry() -> dict:
    """Geometrie je Stab: Dreiecksnetz (Volumenkoerper) oder Achslinie (Zugstaebe)."""
    data = load_model()
    out = {}
    for m in data["members"]:
        out[str(m["id"])] = {"mesh": m["mesh"]} if "mesh" in m else {"line": m["line"]}
    return out


@app.get("/api/export.xlsx")
def export_xlsx() -> StreamingResponse:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    data = load_model()
    wb = Workbook()
    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="305496")

    def sheet(ws, header, rows, widths):
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill = head_font, head_fill
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"

    ws1 = wb.active
    ws1.title = "Querschnitte"
    stat = {"ok": "ok", "info": "Hinweis", "warn": "Warnung", "fail": "Widerspruch"}
    sheet(ws1, ["Querschnitt", "Anzahl", "Flaeche nominal [mm2]", "Summe Laenge [m]", "Volumen [m3]",
                "Material", "Sicherheit", "Pruefung", "Begruendung"],
          [[s["name"], s["count"], s["area_nom_mm2"], s["length_m"], s["volume_m3"], s["material"],
            s["material_confidence"], stat[s["material_status"]], s["material_reason"]] for s in data["sections"]],
          [34, 10, 22, 18, 14, 10, 12, 14, 90])
    n = len(data["sections"]) + 1
    ws1.append(["SUMME", f"=SUM(B2:B{n})", None, f"=SUM(D2:D{n})", f"=SUM(E2:E{n})"])
    for c in ws1[n + 1]:
        c.font = Font(bold=True)

    ws3 = wb.create_sheet("Materialien")
    sheet(ws3, ["Material", "Bezeichnung", "Dichte [kg/m3]", "Anzahl Staebe", "Summe Laenge [m]", "Volumen [m3]",
                "Masse [kg]", "Querschnitte", "Pruefung", "Hinweis"],
          [[r["name"], r["label"], r["density_kg_m3"], r["count"], r["length_m"], r["volume_m3"], r["mass_kg"],
            ", ".join(f"{s['name']} ({s['count']})" for s in r["sections"]), stat[r["status"]],
            r["unused_reason"] or ""] for r in data["materials"]],
          [10, 30, 14, 14, 16, 14, 12, 80, 14, 90])

    ws2 = wb.create_sheet("Staebe", 1)
    sheet(ws2, ["Stab-ID", "Querschnitt", "Querschnitt Ende", "Typ", "Laenge [m]", "Flaeche nominal [mm2]",
                "Flaeche mittel [mm2]", "Volumen [m3]", "Volumen-Quelle", "Material", "Material im Modell",
                "Masse [kg]", "Pruefung", "Hinweis"],
          [[m["id"], m["section"], m["section_end"], m["type"], m["length_m"], m["area_nom_mm2"],
            m["area_mean_mm2"], m["volume_m3"], m["volume_source"], m["material"], m["material_model"],
            m["mass_kg"], stat[m["material_status"]], m["note"]] for m in data["members"]],
          [9, 32, 22, 10, 12, 20, 20, 14, 14, 10, 18, 12, 14, 60])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="Staebe_Material_Stockerau.xlsx"'},
    )
