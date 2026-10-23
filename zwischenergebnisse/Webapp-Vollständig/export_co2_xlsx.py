"""Exportiert kg CO2e je Stab und Lebenszyklus-Modul (GWP-total, live aus der OEKOBAUDAT) als Excel.

Aufruf:  python export_co2_xlsx.py [--steel alternative] [--concrete alternative]
"""
import argparse
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.carbon import compute, module_totals
from app.oekobaudat import MODULE_ORDER

OUT = Path(__file__).parent / "CO2e_Staebe_Stockerau.xlsx"
HEAD_FILL = PatternFill("solid", fgColor="305496")
SUM_FILL = PatternFill("solid", fgColor="D9E1F2")
HINTS = [
    "Indikator: GWP-total in kg CO2e. Werte live aus https://oekobaudat.de (Prozess-Datensaetze, JSON).",
    "Bezugseinheiten: Baustahl 1000 kg, Beton 1 m3, Betonstahl 1 kg (Blatt 'Datensaetze').",
    "Stahl: Masse = Volumen x Rohdichte 7850 kg/m3. Beton: Volumen x GWP je m3.",
    "Bewehrung der Fundamente pauschal 85 kg Betonstahl je m3 Beton.",
    "RO 139.7x8.0: im Modell massiv, hier als Hohlquerschnitt A x L gerechnet. RD 12: Volumen = A x L.",
    "Leere Zellen: Modul im Datensatz nicht deklariert (z. B. A4/A5/B1 bei Stahl) und somit nicht in der Summe.",
    "Modul D (Gutschriften/Lasten ausserhalb der Systemgrenze) nicht mit A-C aufsummieren.",
    "Baustahl: Szenario laut Datensatz-Default (S3, EU-Durchschnitt 88 % Recycling, 11 % Wiederverwendung).",
]


def _header(ws):
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 34


def _sum_row(ws):
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
        c.fill = SUM_FILL


def _widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--steel", choices=["standard", "alternative"], default="standard")
    ap.add_argument("--concrete", choices=["standard", "alternative"], default="standard")
    args = ap.parse_args()
    res = compute(steel=args.steel, concrete=args.concrete)
    mods = [m for m in MODULE_ORDER if any(m in e.gwp for e in res["epds"])]
    n = len(mods)

    wb = Workbook()
    ws = wb.active
    ws.title = "Staebe"
    ws.append(["Stab-ID", "Querschnitt", "Material", "Laenge [m]", "Volumen [m3]", "Masse [kg]",
               "davon Bewehrung [kg]"] + [f"{m} [kg CO2e]" for m in mods] + ["Hinweis"])
    _header(ws)
    for m in res["members"]:
        ws.append([m["id"], m["section"], m["material"], m["length_m"], m["volume_m3"],
                   round(m["mass_kg"], 3), round(m["rebar_kg"], 3)]
                  + [m["gwp"].get(mod) for mod in mods] + [m["note"]])
    last = ws.max_row
    ws.append(["SUMME", None, None] + [f"=SUM({get_column_letter(c)}2:{get_column_letter(c)}{last})"
                                       for c in range(4, 8 + n)])
    _sum_row(ws)
    for r in range(2, ws.max_row + 1):
        for col in range(4, 8 + n):
            ws.cell(r, col).number_format = "#,##0.000" if col < 8 else "#,##0.00"
    _widths(ws, [9, 32, 10, 11, 12, 12, 14] + [14] * n + [60])
    ws.freeze_panes = "C2"

    ws3 = wb.create_sheet("Je Querschnitt")
    ws3.append(["Querschnitt", "Material", "Anzahl"] + [f"{m} [kg CO2e]" for m in mods])
    _header(ws3)
    groups: dict[tuple, dict] = {}
    for m in res["members"]:
        g = groups.setdefault((m["section"], m["material"]), {"n": 0, "gwp": {}})
        g["n"] += 1
        for mod, v in m["gwp"].items():
            if v is not None:
                g["gwp"][mod] = g["gwp"].get(mod, 0.0) + v
    for (sec, mat), g in sorted(groups.items(), key=lambda kv: -kv[1]["gwp"].get("A1-A3", 0)):
        ws3.append([sec, mat, g["n"]] + [g["gwp"].get(mod) for mod in mods])
    tot = module_totals(res)
    ws3.append(["GESAMT", None, len(res["members"])] + [tot.get(mod) for mod in mods])
    _sum_row(ws3)
    for r in range(2, ws3.max_row + 1):
        for col in range(4, 4 + n):
            ws3.cell(r, col).number_format = "#,##0.0"
    _widths(ws3, [34, 10, 8] + [14] * n)
    ws3.freeze_panes = "A2"

    ws2 = wb.create_sheet("Datensaetze")
    ws2.append(["Rolle", "Datensatz", "UUID", "Bezugseinheit", "Bezugsmenge", "Rohdichte [kg/m3]",
                "Abgerufen"] + [f"{m} je Bezug" for m in mods] + ["Szenarien", "Quelle (URL)", "Hinweise"])
    _header(ws2)
    for role, e in zip(["Baustahl", "Beton", "Bewehrung (Betonstahl)"], res["epds"]):
        ws2.append([role, e.name, e.uuid, e.ref_unit, e.ref_amount, e.density_kg_m3, e.fetched]
                   + [e.gwp.get(m) for m in mods]
                   + ["; ".join(f"{k}: {v}" for k, v in e.scenarios.items()), e.url, "; ".join(e.notes)])
    _widths(ws2, [22, 55, 38, 12, 12, 14, 20] + [12] * n + [40, 60, 40])

    ws4 = wb.create_sheet("Hinweise")
    ws4.append([f"Gewaehlt: Baustahl = {args.steel}, Beton = {args.concrete}, Betonstahl = standard."])
    for line in HINTS:
        ws4.append([line])
    ws4.column_dimensions["A"].width = 140

    wb.save(OUT)
    print("gespeichert:", OUT)
    print("Staebe:", len(res["members"]))
    for mod, v in tot.items():
        print(f"  {mod:6s} {v:12.1f} kg CO2e")


if __name__ == "__main__":
    main()
