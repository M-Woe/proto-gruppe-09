"""Web-App: 3D-Ansicht des Rhino-Modells + Staebe-Tabelle (Auswahl im Modell markiert die Zeile).

Start:  python -m uvicorn app.main:app --reload      ->  http://127.0.0.1:8000
"""
from __future__ import annotations

import io
from pathlib import Path

from dataclasses import asdict

from fastapi import FastAPI, HTTPException
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles

from .carbon import DATASETS, compute
from .model import load_model
from .oekobaudat import MODULE_ORDER, OekobaudatError, load_epd

STATIC = Path(__file__).parent / "static"

app = FastAPI(title="Stockerau Staebe-Viewer")
app.mount("/static", StaticFiles(directory=STATIC), name="static")


@app.get("/")
def index() -> FileResponse:
    return FileResponse(STATIC / "index.html")


@app.get("/api/members")
def members() -> dict:
    """Tabellendaten ohne Geometrie (leicht)."""
    data = load_model()
    light = [{k: v for k, v in m.items() if k not in ("mesh", "line")} for m in data["members"]]
    return {
        "source": data["source"],
        "member_count": data["member_count"],
        "sections": data["sections"],
        "members": light,
    }


@app.get("/api/geometry")
def geometry() -> dict:
    """Geometrie je Stab: Dreiecksnetz (Volumenkoerper) oder Achslinie (Zugstaebe)."""
    data = load_model()
    out = {}
    for m in data["members"]:
        out[str(m["id"])] = {"mesh": m["mesh"]} if "mesh" in m else {"line": m["line"]}
    return out


@app.get("/api/carbon")
def carbon(steel: str = "standard", concrete: str = "standard", refresh: bool = False) -> dict:
    """kg CO2e je Stab und Modul (GWP-total), live aus der OEKOBAUDAT.

    steel / concrete: 'standard' oder 'alternative'. refresh=true holt die Datensaetze erneut ab
    (sonst werden sie pro Serverlauf einmal geladen und zwischengespeichert).
    """
    if steel not in DATASETS["steel"] or concrete not in DATASETS["concrete"]:
        raise HTTPException(400, "steel und concrete muessen 'standard' oder 'alternative' sein")
    if refresh:
        load_epd.cache_clear()
    try:
        res = compute(steel=steel, concrete=concrete)
    except OekobaudatError as exc:
        raise HTTPException(502, str(exc)) from exc
    epds = []
    for role, e in zip(("Baustahl", "Beton", "Bewehrung (Betonstahl)"), res["epds"]):
        d = asdict(e)
        d.update(role=role, modules=e.modules)
        epds.append(d)
    members = [{
        "id": m["id"], "section": m["section"], "material": m["material"], "length_m": m["length_m"],
        "volume_m3": m["volume_m3"], "mass_kg": m["mass_kg"], "rebar_kg": m["rebar_kg"],
        "gwp": m["gwp"], "note": m["note"],
        "parts": [{k: p[k] for k in ("label", "dataset", "amount", "unit", "gwp")} for p in m["parts"]],
    } for m in res["members"]]
    modules = [k for k in MODULE_ORDER if any(k in m["gwp"] for m in members)]
    return {"options": res["options"], "modules": modules, "epds": epds, "members": members}


@app.get("/api/export.xlsx")
def export_xlsx() -> StreamingResponse:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    data = load_model()
    wb = Workbook()
    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="305496")

    def sheet(ws, header, rows, widths):
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill = head_font, head_fill
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"

    ws1 = wb.active
    ws1.title = "Querschnitte"
    sheet(ws1, ["Querschnitt", "Anzahl", "Flaeche nominal [mm2]", "Summe Laenge [m]", "Volumen [m3]"],
          [[s["name"], s["count"], s["area_nom_mm2"], s["length_m"], s["volume_m3"]] for s in data["sections"]],
          [34, 10, 22, 18, 14])
    n = len(data["sections"]) + 1
    ws1.append(["SUMME", f"=SUM(B2:B{n})", None, f"=SUM(D2:D{n})", f"=SUM(E2:E{n})"])
    for c in ws1[n + 1]:
        c.font = Font(bold=True)

    ws2 = wb.create_sheet("Staebe")
    sheet(ws2, ["Stab-ID", "Querschnitt", "Querschnitt Ende", "Typ", "Laenge [m]", "Flaeche nominal [mm2]",
                "Flaeche mittel [mm2]", "Volumen [m3]", "Volumen-Quelle", "Hinweis"],
          [[m["id"], m["section"], m["section_end"], m["type"], m["length_m"], m["area_nom_mm2"],
            m["area_mean_mm2"], m["volume_m3"], m["volume_source"], m["note"]] for m in data["members"]],
          [9, 32, 22, 10, 12, 20, 20, 14, 14, 60])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="Staebe_Stockerau.xlsx"'},
    )
