"""Exportiert die Staebe-Auswertung (Stockerau, RFEM-Export) als Excel-Tabelle."""
import math
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUT = Path(__file__).parent / "Data" / "Staebe_Auswertung_Stockerau.xlsx"

# (Querschnitt, Anzahl, Querschnittsflaeche [mm2], Achsenlaenge [m], Volumen [m3], Hinweis)
# Volumen aus Rhino-Geometrie (Members-Layer); RD 12 hat nur Achslinien -> berechnet.
FLACH = [
    ("Rectangle 500/500", 47, 500 * 500, 47.023, 11.73823, "Volumenkoerper"),
    ("RO 139.7x8.0 | EN 10210-2:2006", 12, math.pi / 4 * (139.7**2 - 123.7**2), 41.070, 0.62946,
     "Im Modell massiver Zylinder (Volumen = Vollquerschnitt)"),
    ("Flat Bar 275/20", 41, 275 * 20, 26.454, 0.13275, "Volumenkoerper"),
    ("Flat Bar 275/12", 70, 275 * 12, 42.638, 0.12853, "Volumenkoerper"),
    ("Flat Bar 225/20", 23, 225 * 20, 19.443, 0.07748, "Volumenkoerper"),
    ("Flat Bar 225/12", 35, 225 * 12, 29.587, 0.07074, "Volumenkoerper"),
    ("Flat Bar 175/20", 23, 175 * 20, 19.285, 0.05757, "Volumenkoerper"),
    ("Flat Bar 175/12", 35, 175 * 12, 29.587, 0.05299, "Volumenkoerper"),
    ("Flat Bar 237.5/12", 35, 237.5 * 12, 13.639, 0.03568, "Volumenkoerper"),
    ("Flat Bar 237.5/20", 18, 237.5 * 20, 7.014, 0.03058, "Volumenkoerper"),
    ("Flat Bar 200/12", 35, 200 * 12, 13.649, 0.02956, "Volumenkoerper"),
    ("Flat Bar 200/20", 18, 200 * 20, 7.019, 0.02534, "Volumenkoerper"),
    ("Flat Bar 162.5/12", 35, 162.5 * 12, 13.639, 0.02340, "Volumenkoerper"),
    ("Flat Bar 162.5/20", 18, 162.5 * 20, 6.842, 0.01957, "Volumenkoerper"),
    ("Flat Bar 33.3/20", 116, 33.3 * 20, 2.320, 0.00151, "Volumenkoerper"),
]
RD12_LEN = 367.168
RD12_AREA = math.pi / 4 * 12**2
RD12 = ("RD 12 (Rundstab)", 175, RD12_AREA, RD12_LEN, RD12_AREA * 1e-6 * RD12_LEN,
        "Nur Achslinien im Modell; Volumen = A x L berechnet")

ROWS = FLACH + [RD12]


def main() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Staebe"

    header = ["Querschnitt", "Anzahl Staebe", "Querschnittsflaeche [mm2]",
              "Querschnittsflaeche [m2]", "Summe Achsenlaenge [m]",
              "Volumen gesamt [m3]", "Volumen je Stab [m3]", "Hinweis"]
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for i, (name, n, area_mm2, length, vol, note) in enumerate(ROWS, start=2):
        ws.append([name, n, round(area_mm2, 1), f"=C{i}/1000000", length, vol, f"=F{i}/B{i}", note])

    last = len(ROWS) + 1
    tot = last + 1
    ws.append(["SUMME", f"=SUM(B2:B{last})", None, None, f"=SUM(E2:E{last})", f"=SUM(F2:F{last})", None,
               "Summe inkl. berechnetem Volumen RD 12"])
    for c in ws[tot]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="D9E1F2")

    for r in range(2, tot + 1):
        ws.cell(r, 2).number_format = "0"
        ws.cell(r, 3).number_format = "#,##0.0"
        ws.cell(r, 4).number_format = "0.000000"
        ws.cell(r, 5).number_format = "0.00"
        ws.cell(r, 6).number_format = "0.0000"
        ws.cell(r, 7).number_format = "0.000000"

    for col, w in enumerate([34, 14, 18, 18, 18, 16, 16, 55], start=1):
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"

    ws2 = wb.create_sheet("Hinweise")
    for line in [
        "Quelle: Data/Stockerau Rhino Export from RFEM.3dm (Einheit: Meter)",
        "Layer: 2__B+G::03_3D::RFEM_Export_260915::Members bzw. ::MemberAxes",
        "Stabanzahl: 561 Volumenkoerper (Brep) + 175 RD-12-Achslinien = 736 Staebe",
        "Querschnittsflaeche = nominal aus Profilbezeichnung (Flachstahl: Hoehe x Dicke).",
        "Modellierte Flachstahl-Koerper sind teils kleiner (variable Hoehe), Volumen stammt aus der Geometrie.",
        "RO 139.7x8.0: Flaeche als Rohrquerschnitt; im Modell aber massiv -> Volumen entspricht Vollzylinder.",
    ]:
        ws2.append([line])
    ws2.column_dimensions["A"].width = 110

    wb.save(OUT)
    print("gespeichert:", OUT)


if __name__ == "__main__":
    main()
