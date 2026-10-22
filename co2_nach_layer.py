"""
Gruppiere CO2-Daten nach Layer (eine Zeile pro Profiltyp).
"""
import json, csv
from collections import defaultdict

with open("data/staebe_mit_co2.json") as f:
    data = json.load(f)

# Gruppieren
by_layer = defaultdict(lambda: {
    "n":0, "V":0, "M":0, "L":0,
    "A1A3":0, "C":0, "D":0,
    "bew_kg":0, "bew_A1A3":0
})

for r in data:
    b = by_layer[r["layer"]]
    b["n"] += 1
    b["V"] += r["volumen_m3"]
    b["M"] += r["masse_kg"]
    b["L"] += r["laenge_m"]
    b["A1A3"] += r["co2_A1A3_kg"]
    b["C"] += r["co2_C_sum_kg"]
    b["D"] += r["co2_D_kg"]
    b["bew_kg"] += r["bewehrung_kg"]
    b["bew_A1A3"] += r["bewehrung_co2_A1A3_kg"]

rows = []
for ln in sorted(by_layer):
    b = by_layer[ln]
    rows.append({
        "Layer": ln,
        "Anzahl": b["n"],
        "Laenge_m": round(b["L"], 3),
        "Volumen_m3": round(b["V"], 6),
        "Masse_kg": round(b["M"], 2) if b["M"] > 0 else 0,
        "CO2_A1A3_gesamt_kg": round(b["A1A3"], 4),
        "CO2_A1A3_pro_Stab_kg": round(b["A1A3"] / b["n"], 6),
        "CO2_C_gesamt_kg": round(b["C"], 4),
        "CO2_D_gesamt_kg": round(b["D"], 4),
        "Bewehrung_kg": round(b["bew_kg"], 2) if b["bew_kg"] > 0 else 0,
        "Bewehrung_CO2_A1A3_kg": round(b["bew_A1A3"], 4) if b["bew_A1A3"] > 0 else 0,
    })

# Summenzeile
s = {k: sum(r[k] for r in rows if isinstance(r[k], (int, float))) for k in rows[0].keys() if k != "Layer"}
rows.append({
    "Layer": "SUMME", "Anzahl": int(s["Anzahl"]),
    "Laenge_m": round(s["Laenge_m"], 3),
    "Volumen_m3": round(s["Volumen_m3"], 6),
    "Masse_kg": round(s["Masse_kg"], 2) if s["Masse_kg"] > 0 else 0,
    "CO2_A1A3_gesamt_kg": round(s["CO2_A1A3_gesamt_kg"], 4),
    "CO2_A1A3_pro_Stab_kg": 0,
    "CO2_C_gesamt_kg": round(s["CO2_C_gesamt_kg"], 4),
    "CO2_D_gesamt_kg": round(s["CO2_D_gesamt_kg"], 4),
    "Bewehrung_kg": round(s["Bewehrung_kg"], 2) if s["Bewehrung_kg"] > 0 else 0,
    "Bewehrung_CO2_A1A3_kg": round(s["Bewehrung_CO2_A1A3_kg"], 4) if s["Bewehrung_CO2_A1A3_kg"] > 0 else 0,
})

# Ausgabe
print(f"{'Layer':35s} {'Anz':>4s} {'CO2/Stab':>10s} {'CO2 ges.':>10s} {'Bew. A1A3':>10s}")
print("-"*75)
for r in rows[:-1]:
    print(f"  {r['Layer']:35s} {r['Anzahl']:4d} {r['CO2_A1A3_pro_Stab_kg']:10.6f} {r['CO2_A1A3_gesamt_kg']:10.4f} {r['Bewehrung_CO2_A1A3_kg']:10.4f}")
r = rows[-1]
print(f"  {r['Layer']:35s} {r['Anzahl']:4d} {'-':>10s} {r['CO2_A1A3_gesamt_kg']:10.4f} {r['Bewehrung_CO2_A1A3_kg']:10.4f}")
print(f"\nGesamt CO2 A1-A3: {s['CO2_A1A3_gesamt_kg']+s['Bewehrung_CO2_A1A3_kg']:.2f} kg = {(s['CO2_A1A3_gesamt_kg']+s['Bewehrung_CO2_A1A3_kg'])/1000:.3f} t")

# CSV
csv_path = "data/co2_nach_layer.csv"
with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
    w = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
    w.writeheader()
    w.writerows(rows)
print(f"CSV: {csv_path} ({len(rows)} Zeilen)")

# JSON
json_path = "data/co2_nach_layer.json"
with open(json_path, "w", encoding="utf-8") as f:
    json.dump(rows, f, indent=2, ensure_ascii=False)
print(f"JSON: {json_path}")

# Excel-Sheet
import openpyxl
from openpyxl.utils import get_column_letter

xlsx_path = "data/Aufgabe 1 Excel-Liste.xlsx"
wb = openpyxl.load_workbook(xlsx_path)
if "co2_nach_layer" in wb.sheetnames:
    del wb["co2_nach_layer"]
ws = wb.create_sheet("co2_nach_layer")
headers = list(rows[0].keys())
num = set(headers) - {"Layer"}
for c, h in enumerate(headers, 1):
    ws.cell(1, c, h)
for ri, r in enumerate(rows, 2):
    for ci, h in enumerate(headers, 1):
        v = r.get(h, "")
        if h in num and isinstance(v, (int, float)) and v != 0:
            ws.cell(ri, ci, float(v))
        else:
            ws.cell(ri, ci, str(v))
for ci, h in enumerate(headers, 1):
    ws.column_dimensions[get_column_letter(ci)].width = max(12, len(h)+2)
    if h in num:
        for ri in range(2, len(rows)+2):
            ws.cell(ri, ci).number_format = "0.000000"
try:
    wb.save(xlsx_path)
    print(f"Excel: Sheet 'co2_nach_layer' in {xlsx_path}")
except PermissionError:
    alt = "data/Aufgabe 1 Excel-Liste (aggregiert).xlsx"
    wb.save(alt)
    print(f"Excel war gesperrt -> gespeichert als: {alt}")