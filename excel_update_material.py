"""
Fuege Material-Zuordnung als Sheet in die Excel-Tabelle ein.
"""
import openpyxl, json
from openpyxl.utils import get_column_letter

xlsx_path = "data/Aufgabe 1 Excel-Liste.xlsx"
with open("data/staebe_mit_material.json") as f:
    data = json.load(f)

wb = openpyxl.load_workbook(xlsx_path)
if "staebe_mit_material" in wb.sheetnames:
    del wb["staebe_mit_material"]

ws = wb.create_sheet("staebe_mit_material")
headers = list(data[0].keys())
numeric_cols = {"laenge_m", "volumen_m3"}

for c, h in enumerate(headers, 1):
    ws.cell(1, c, h)
for ri, row in enumerate(data, 2):
    for ci, h in enumerate(headers, 1):
        v = row.get(h, "")
        if h in numeric_cols and v != "":
            try:
                ws.cell(ri, ci, float(v))
                continue
            except:
                pass
        ws.cell(ri, ci, str(v))

for ci, h in enumerate(headers, 1):
    ws.column_dimensions[get_column_letter(ci)].width = max(10, len(h)+2)
for ci, h in enumerate(headers, 1):
    if h in numeric_cols:
        for ri in range(2, len(data)+2):
            ws.cell(ri, ci).number_format = "0.00000000"

wb.save(xlsx_path)
print(f"Excel aktualisiert: {xlsx_path}")
print(f"  Sheet 'staebe_mit_material' ({len(data)} Zeilen)")