"""
Etappe 1: Alle aus dem Rhino-Modell ausgelesenen Daten als Excel-Datei exportieren.

Speicherort: <Projekt>/output/excel/etappe1_modell_export.xlsx
(Der Ordner wird bei Bedarf automatisch angelegt.)

Sheets:
- Stabliste:        eine Zeile je Stab (RFEM-BeamId), wie data/etappe1_staebe_liste.csv
- Achsen_Rohdaten:  alle LineCurve-Objekte mit BeamId inkl. aller Benutzerdaten
- Brep_Rohdaten:    alle Brep-Objekte mit BeamId inkl. aller Benutzerdaten
- Objekttypen:      Anzahl je Geometrietyp im gesamten Modell
- Layer:            alle Layer mit Objektzahlen je Typ
- Zusammenfassung:  Stabanzahl je Typ und Layer, Gesamtlänge
"""
import math
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl
import rhino3dm
from openpyxl.utils import get_column_letter

from stab_liste_etappe1 import CSV_FIELDS, MODEL_PATH, build_stab_liste

ROOT = Path(__file__).resolve().parent
OUT_DIR = ROOT / "output" / "excel"
OUT_FILE = OUT_DIR / "etappe1_modell_export.xlsx"


def _write_sheet(ws, headers, rows):
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    for ci, h in enumerate(headers, 1):
        width = max([len(str(h))] + [len(str(r.get(h, ""))) for r in rows[:500]])
        ws.column_dimensions[get_column_letter(ci)].width = min(max(10, width + 2), 60)
    ws.freeze_panes = "A2"


def _raw_rows(model_path: Path, want_type: str) -> tuple[list, list]:
    """Liefert (Zeilen, Benutzerdaten-Schlüssel) aller Objekte mit BeamId vom Typ want_type."""
    model = rhino3dm.File3dm.Read(str(model_path))
    layer_names = {layer.Index: layer.Name for layer in model.Layers}
    rows, keys = [], []
    for idx, obj in enumerate(model.Objects):
        geom = obj.Geometry
        if geom is None or type(geom).__name__ != want_type:
            continue
        ud = dict(obj.Attributes.GetUserStrings())
        if "BG_RFEM_Exchange_BeamId_bg0246" not in ud:
            continue
        row = {"obj_index": idx, "uuid": str(obj.Attributes.Id),
               "layer": layer_names.get(obj.Attributes.LayerIndex, ""),
               "geometrietyp": want_type, "objekt_name": obj.Attributes.Name or ""}
        if want_type == "LineCurve":
            p0, p1 = geom.PointAtStart, geom.PointAtEnd
            row.update({"start_x": p0.X, "start_y": p0.Y, "start_z": p0.Z,
                        "ende_x": p1.X, "ende_y": p1.Y, "ende_z": p1.Z,
                        "laenge_m": math.dist((p0.X, p0.Y, p0.Z), (p1.X, p1.Y, p1.Z))})
        else:
            bb = geom.GetBoundingBox()
            row.update({"bb_min_x": bb.Min.X, "bb_min_y": bb.Min.Y, "bb_min_z": bb.Min.Z,
                        "bb_max_x": bb.Max.X, "bb_max_y": bb.Max.Y, "bb_max_z": bb.Max.Z,
                        "ist_solid": bool(geom.IsSolid)})
        for k, v in ud.items():
            row[k] = v
            if k not in keys:
                keys.append(k)
        rows.append(row)
    return rows, keys


def _model_overview(model_path: Path) -> tuple[list, list]:
    """Objekttypen und Layer-Übersicht des gesamten Modells."""
    model = rhino3dm.File3dm.Read(str(model_path))
    layer_names = {layer.Index: layer.Name for layer in model.Layers}
    type_counts = Counter(type(o.Geometry).__name__ if o.Geometry else "None" for o in model.Objects)
    per_layer = defaultdict(Counter)
    for o in model.Objects:
        per_layer[layer_names.get(o.Attributes.LayerIndex, "")][type(o.Geometry).__name__ if o.Geometry else "None"] += 1
    type_rows = [{"geometrietyp": t, "anzahl": n} for t, n in type_counts.most_common()]
    layer_rows = [{"layer": ln, "gesamt": sum(c.values()), **{t: c[t] for t in sorted(c)}}
                  for ln, c in sorted(per_layer.items())]
    return type_rows, layer_rows


def _summary_rows(stabe: list) -> list:
    """Stabanzahl und Länge je Typ und Layer."""
    by_key = defaultdict(lambda: {"anzahl": 0, "laenge_m": 0.0})
    for s in stabe:
        b = by_key[(s["rfem_typ"], s["layer"])]
        b["anzahl"] += 1
        b["laenge_m"] += s["laenge_m"]
    rows = [{"rfem_typ": t, "layer": ln, "anzahl": b["anzahl"], "laenge_m": round(b["laenge_m"], 4)}
            for (t, ln), b in sorted(by_key.items())]
    rows.append({"rfem_typ": "SUMME", "layer": "",
                 "anzahl": len(stabe), "laenge_m": round(sum(s["laenge_m"] for s in stabe), 4)})
    return rows


def export(out_file: Path = OUT_FILE) -> dict:
    out_file.parent.mkdir(parents=True, exist_ok=True)

    stabe = build_stab_liste(MODEL_PATH)
    achsen, achsen_keys = _raw_rows(MODEL_PATH, "LineCurve")
    breps, brep_keys = _raw_rows(MODEL_PATH, "Brep")
    typen, layer_rows = _model_overview(MODEL_PATH)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    _write_sheet(wb.create_sheet("Stabliste"), CSV_FIELDS, stabe)
    _write_sheet(wb.create_sheet("Achsen_Rohdaten"),
                 ["obj_index", "uuid", "layer", "geometrietyp", "objekt_name",
                  "start_x", "start_y", "start_z", "ende_x", "ende_y", "ende_z", "laenge_m"] + achsen_keys,
                 achsen)
    _write_sheet(wb.create_sheet("Brep_Rohdaten"),
                 ["obj_index", "uuid", "layer", "geometrietyp", "objekt_name",
                  "bb_min_x", "bb_min_y", "bb_min_z", "bb_max_x", "bb_max_y", "bb_max_z",
                  "ist_solid"] + brep_keys,
                 breps)
    _write_sheet(wb.create_sheet("Objekttypen"), ["geometrietyp", "anzahl"], typen)
    layer_headers = ["layer", "gesamt"] + sorted({k for r in layer_rows for k in r} - {"layer", "gesamt"})
    _write_sheet(wb.create_sheet("Layer"), layer_headers, layer_rows)
    _write_sheet(wb.create_sheet("Zusammenfassung"),
                 ["rfem_typ", "layer", "anzahl", "laenge_m"], _summary_rows(stabe))

    wb.save(out_file)
    return {"datei": str(out_file), "stabe": len(stabe), "achsen": len(achsen),
            "breps": len(breps), "sheets": wb.sheetnames}


if __name__ == "__main__":
    result = export()
    print(f"Excel geschrieben: {result['datei']}")
    print(f"Stäbe: {result['stabe']} | Achsen-Rohzeilen: {result['achsen']} | Brep-Rohzeilen: {result['breps']}")
    print(f"Sheets: {result['sheets']}")
