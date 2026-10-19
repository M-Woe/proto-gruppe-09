"""Lies die Excel-Liste (Aufgabe 1) und zeige den Inhalt."""
import openpyxl

wb = openpyxl.load_workbook("data/Aufgabe 1 Excel-Liste.xlsx", data_only=True)
print("Sheet names:", wb.sheetnames)

for sname in wb.sheetnames:
    ws = wb[sname]
    print(f"\n=== Sheet: {sname} ({ws.max_row} Zeilen x {ws.max_column} Spalten) ===")
    for r in range(1, min(ws.max_row + 1, 60)):
        row_data = [str(ws.cell(r, c).value)[:80] for c in range(1, ws.max_column + 1)]
        print(f"  R{r:3d}: {row_data}")
    if ws.max_row > 60:
        print(f"  ... ({ws.max_row - 60} weitere Zeilen)")