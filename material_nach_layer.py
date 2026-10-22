"""
Gruppiere Material-Daten nach Layer.
"""
import json, csv
from collections import defaultdict

with open("data/staebe_mit_material.json") as f:
    data = json.load(f)

# Gruppieren
by = defaultdict(lambda: {"n":0, "L":0, "V":0, "material":"", "kategorie":"", "norm":""})
for r in data:
    b = by[r["layer"]]
    b["n"] += 1
    b["L"] += r["laenge_m"]
    b["V"] += r["volumen_m3"]
    b["material"] = r["material"]
    b["kategorie"] = r["material_category"]
    b["norm"] = r["material_standard"]

rows = []
for ln in sorted(by):
    b = by[ln]
    rows.append({
        "Layer": ln,
        "Anzahl": b["n"],
        "Material": b["material"],
        "Kategorie": b["kategorie"],
        "Norm": b["norm"],
        "Laenge_m": round(b["L"], 3),
        "Volumen_m3": round(b["V"], 6),
    })

# Summenzeile
rows.append({
    "Layer": "SUMME",
    "Anzahl": sum(r["Anzahl"] for r in rows),
    "Material": "", "Kategorie": "", "Norm": "",
    "Laenge_m": round(sum(r["Laenge_m"] for r in rows), 3),
    "Volumen_m3": round(sum(r["Volumen_m3"] for r in rows), 6),
})

# Ausgabe
print(f"{'Layer':35s} {'Anz':>4s} {'Material':25s} {'Laenge_m':>8s} {'Vol_m3':>8s}")
print("-"*85)
for r in rows:
    print(f"  {r['Layer']:35s} {r['Anzahl']:4d} {r['Material']:25s} {r['Laenge_m']:8.3f} {r['Volumen_m3']:8.6f}")

# CSV
csv_path = "data/material_nach_layer.csv"
with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
    w = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
    w.writeheader(); w.writerows(rows)
print(f"\nCSV: {csv_path}")

# JSON
json_path = "data/material_nach_layer.json"
with open(json_path, "w", encoding="utf-8") as f:
    json.dump(rows, f, indent=2, ensure_ascii=False)
print(f"JSON: {json_path}")

# Excel
import openpyxl
from openpyxl.utils import get_column_letter

xlsx = "data/Aufgabe 1 Excel-Liste.xlsx"
wb = openpyxl.load_workbook(xlsx)
if "material_nach_layer" in wb.sheetnames:
    del wb["material_nach_layer"]
ws = wb.create_sheet("material_nach_layer")
hdrs = list(rows[0].keys())
num = {"Anzahl","Laenge_m","Volumen_m3"}
for c, h in enumerate(hdrs, 1):
    ws.cell(1, c, h)
for ri, r in enumerate(rows, 2):
    for ci, h in enumerate(hdrs, 1):
        v = r.get(h, "")
        if h in num and isinstance(v, (int, float)):
            ws.cell(ri, ci, float(v))
        else:
            ws.cell(ri, ci, str(v))
for ci, h in enumerate(hdrs, 1):
    ws.column_dimensions[get_column_letter(ci)].width = max(12, len(h)+2)
try:
    wb.save(xlsx)
    print(f"Excel: Sheet 'material_nach_layer' in {xlsx}")
except PermissionError:
    alt = "data/Aufgabe 1 Excel-Liste (aggregiert).xlsx"
    wb.save(alt)
    print(f"Excel gesperrt -> {alt}")